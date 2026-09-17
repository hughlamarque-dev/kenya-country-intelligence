"""Read downloaded workbooks; never modifies or authors Excel files."""
import json,sys,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import load_workbook
folder=Path(sys.argv[1]);expected=json.loads((folder/'expected.json').read_text())
def rows(book,sheet):
    w=book[sheet];headers=[c.value for c in w[5]]
    return [dict(zip(headers,r)) for r in w.iter_rows(min_row=6,values_only=True)]
books={name:load_workbook(folder/(name+'.xlsx')) for name in ['filtered','routes','shortlist','individual','edge-cases']}
for name,w in books.items():
    with zipfile.ZipFile(folder/(name+'.xlsx')) as z:
        assert z.testzip() is None
        for p in z.namelist():
            if p.endswith(('.xml','.rels')):ET.fromstring(z.read(p))
    assert len(w.worksheets)==7
    for sh in w:
        assert sh.freeze_panes=='B6'
        assert sh.auto_filter.ref
        assert not any(c.data_type=='f' for row in sh for c in row),'Source text must not become formulas'
        assert not any(c.value in ['#REF!','#DIV/0!','#VALUE!'] for row in sh for c in row)
    source_ids={r['Source ID'] for r in rows(w,'Sources')}
    for sh in ['Partners','Calls and routes','Projects','Funding evidence','Financial history']:
        for r in rows(w,sh):
            if r.get('Source IDs'):assert set(r['Source IDs'].split('; '))<=source_ids
    for row in w['Sources'].iter_rows(min_row=6):
        if row[3].value:assert row[3].hyperlink.target==row[3].value
assert sorted(r['Partner'] for r in rows(books['filtered'],'Partners'))==sorted(expected['filteredNames'])
assert [r['Partner'] for r in rows(books['routes'],'Partners')]==expected['routeNames']
assert [r['Partner'] for r in rows(books['shortlist'],'Partners')]==[expected['savedName']]
assert len(rows(books['shortlist'],'Projects'))>len(rows(books['filtered'],'Projects'))
assert any(r['Original amount']==37090000000 and r['Currency']=='JPY' for r in rows(books['individual'],'Funding evidence'))
assert all(r['Deadline as reported']=='Not separately recorded — see route and sources' for r in rows(books['routes'],'Calls and routes'))
ed=books['edge-cases'];assert ed['Partners']['A6'].value=='=2+2' and ed['Partners']['A6'].data_type=='s'
assert [r['Original amount'] for r in rows(ed,'Funding evidence')][:3]==[-10,0,None]
assert rows(ed,'Funding evidence')[3]['Period start']=='2020-02'
assert ed['Partners']['B6'].value=='A & B <test>'
assert ed['Sources']['D6'].hyperlink.target=='https://example.org/?x=1&y=2'
print(json.dumps({'passed':True,'rows':{name:{sh:len(rows(w,sh)) for sh in w.sheetnames} for name,w in books.items()},'checks':['valid ZIP and XML','readable XLSX','frozen headers','filters','clickable sources','source reference integrity','original JPY amount','missing deadlines preserved','missing vs zero','partial dates','formula injection treated as text']}))
